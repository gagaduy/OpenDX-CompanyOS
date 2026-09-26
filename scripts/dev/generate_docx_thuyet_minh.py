# SPDX-FileCopyrightText: 2026 OpenDX CompanyOS contributors
# SPDX-License-Identifier: Apache-2.0

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_document():
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("BẢN THUYẾT MINH DỰ ÁN: OPENDX COMPANYOS\n")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(17)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    subtitle = title.add_run("Hệ Điều Hành Doanh Nghiệp Tự Trị (Autonomous Enterprise Operating System)")
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(13)
    subtitle.font.italic = True
    subtitle.font.color.rgb = RGBColor(14, 116, 144)

    doc.add_paragraph("─" * 55).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. GIỚI THIỆU TỔNG QUAN
    h1 = doc.add_heading(level=1)
    run_h1 = h1.add_run("1. GIỚI THIỆU TỔNG QUAN")
    run_h1.font.name = "Arial"
    run_h1.font.size = Pt(13.5)
    run_h1.font.bold = True
    run_h1.font.color.rgb = RGBColor(15, 23, 42)

    items_intro = [
        ("Tên hệ thống", "OpenDX CompanyOS"),
        ("Mô hình hoạt động", "Hệ điều hành hợp nhất giữa Nền tảng Thương mại (Commerce Engine), Lực lượng Nhân sự Số đa tác tử (Agentic Workforce) và Bộ điều phối quy trình thị giác (Visual Workflow Studio / iPaaS)."),
        ("Mục tiêu dự án", "Xây dựng một nền tảng vận hành doanh nghiệp mã nguồn mở, cho phép tích hợp AI vào sâu trong quy trình nghiệp vụ thực tế có tổ chức, có phân quyền, có kiểm soát rủi ro bằng con người (Human-in-the-loop) thay vì các chatbot hỗ trợ đơn lẻ."),
        ("Tiêu chuẩn công nghệ", "Kiến trúc sạch (Clean Architecture), Domain-Driven Design (DDD), Temporal Durable Execution, Keycloak RBAC Security, PostgreSQL Authoritative Store, Vite & React Staff Console."),
        ("Bản quyền", "Mã nguồn mở chuẩn Apache-2.0, đáp ứng các tiêu chí đánh giá phần mềm nguồn mở OLP 2026.")
    ]
    for label, val in items_intro:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"• {label}: "); r.bold = True
        p.add_run(val)

    # 2. PAIN POINTS & SOLUTIONS
    h2 = doc.add_heading(level=1)
    run_h2 = h2.add_run("2. CÁC BÀI TOÁN THỰC TẾ (PAIN POINTS) & GIẢI PHÁP KỸ THUẬT")
    run_h2.font.name = "Arial"
    run_h2.font.size = Pt(13.5)
    run_h2.font.bold = True
    run_h2.font.color.rgb = RGBColor(15, 23, 42)

    base_dir = os.path.abspath("docs/screenshots")

    def add_image_block(img_filename, caption_text):
        img_path = os.path.join(base_dir, img_filename)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_picture(img_path, width=Inches(6.2))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_cap = cap.add_run(caption_text)
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(100, 116, 139)

    # PAIN POINT 1
    doc.add_heading(level=2).add_run("PAIN POINT 1: AI trong doanh nghiệp phần lớn chỉ là các chatbot độc lập, không có khả năng phân rã mục tiêu chiến lược và không thể phối hợp liên phòng ban.")
    p = doc.add_paragraph()
    p.add_run("• Vấn đề thực tế: ").bold = True
    p.add_run("Hầu hết các giải pháp hiện nay chỉ dừng lại ở việc cung cấp một ô chat để nhân viên tự đặt câu hỏi. Khi Ban Giám đốc đưa ra một mục tiêu kinh doanh cấp cao (ví dụ: 'Chuẩn bị chiến dịch ra mắt dòng sản phẩm mới' hoặc 'Rà soát và xử lý dứt điểm các khiếu nại khách hàng tồn đọng'), con người vẫn phải tự bóc tách việc thành từng phần nhỏ rồi copy-paste thủ công sang từng công cụ riêng lẻ. Các phòng ban (Marketing, Kho, Chăm sóc khách hàng, Kế toán) hoạt động rời rạc, gây chậm trễ và đứt gãy mạch thông tin.")

    p = doc.add_paragraph()
    p.add_run("• Giải pháp trong OpenDX CompanyOS:").bold = True
    doc.add_paragraph(
        "1. Cơ chế điều phối thông qua AI CEO: Hệ thống thiết kế một tác tử chỉ huy trung tâm (AI CEO). Khi nhận chỉ thị chiến lược bằng ngôn ngữ tự nhiên từ Ban Quản trị, AI CEO tự động phân tích ngữ cảnh, phân rã mục tiêu (Goal Decomposition) thành các gói công việc cụ thể và phân công trực tiếp xuống các nhân sự số chuyên trách.\n"
        "2. Tổ hợp 10 Nhân sự Số chuyên môn hóa: Chuyên viên Marketing (soạn thảo bài viết, tạo ảnh quảng bá), Chuyên viên CSKH (phân tích mức độ bức xúc, soạn email phản hồi và đề xuất mức bồi thường), Chuyên viên Quản lý Kho (giám sát ngưỡng an toàn hàng tồn, cảnh báo bổ sung hàng), Chuyên viên Catalog & Định giá (quản lý cấu trúc sản phẩm, bảo đảm tính nhất quán giá bán), Chuyên viên Tài chính - Kế toán (tổng hợp doanh thu, tính toán LTV).\n"
        "3. Bảo mật phân quyền độc lập (Keycloak RBAC): Mỗi nhân sự AI được cấp một Service Account và Access Token riêng biệt. Không có tác tử nào được dùng chung quyền hoặc tự ý can thiệp vào dữ liệu ngoài phạm vi được chỉ định."
    )
    add_image_block("04_command_center.png", "[Hình 1] Trung tâm Chỉ huy Tác nghiệp (Command Center) - Giao việc chiến lược cho AI CEO & Live Activity Feed")
    add_image_block("06_digital_employees.png", "[Hình 2] Danh bạ quản lý 10 Nhân sự Số (Digital Employees) với RBAC & phân bổ LLM độc lập")

    # PAIN POINT 2
    doc.add_heading(level=2).add_run("PAIN POINT 2: Quy trình tự động hóa nghiệp vụ thường bị 'code cứng', khó tùy biến theo chính sách thay đổi và dễ đứt gãy khi hệ thống gặp sự cố.")
    p = doc.add_paragraph()
    p.add_run("• Vấn đề thực tế: ").bold = True
    p.add_run("Trong các hệ sinh thái doanh nghiệp truyền thống, logic nghiệp vụ thường bị cài cắm trực tiếp vào mã nguồn backend. Mỗi khi doanh nghiệp thay đổi chính sách (ví dụ: thay đổi hạn mức bồi thường hoặc bổ sung bước gửi email thông báo), lập trình viên phải sửa code và deploy lại. Ngược lại, nếu sử dụng các công cụ automation kịch bản thông thường thì thiếu cơ chế duy trì trạng thái bền bỉ (stateful); khi máy chủ khởi động lại hoặc mạng gián đoạn, quy trình đang chạy dở sẽ bị hủy bỏ hoàn toàn.")

    p = doc.add_paragraph()
    p.add_run("• Giải pháp trong OpenDX CompanyOS:").bold = True
    doc.add_paragraph(
        "1. Visual Business Workflow Studio (iPaaS trực quan): Cung cấp giao diện đồ họa Canvas tương tác hoàn chỉnh cho phép người vận hành tự kéo - thả, cấu hình và nối dây các khối chức năng mà không cần sửa mã nguồn.\n"
        "2. Thư viện khối đa dạng: 5 nhóm khối nghiệp vụ trực quan gồm: Sự kiện kích hoạt (Trigger), Khối Nhân sự AI (Agent Worker), Khối Phân luồng điều kiện (Decision Router), Cổng phê duyệt (Approval Gate) và Khối Hành động (Action).\n"
        "3. Phân luồng chính sách thông minh (Conditional Branching): Hỗ trợ rẽ nhánh logic tự động dựa trên giá trị dữ liệu chạy qua luồng. Ví dụ trong quy trình xử lý khiếu nại (WF-CSKH-RECOVERY): Nếu giá trị bồi thường <= 200.000đ thì tự động đi theo nhánh xanh sinh mã voucher và gửi thư xin lỗi ngay; nếu > 200.000đ thì tự động chuyển hướng sang nhánh vàng chờ quản lý thẩm định.\n"
        "4. Vận hành bền bỉ trên Temporal Durable Execution: Tiến trình quy trình được lưu vết trạng thái từng bước vào cơ sở dữ liệu. Ngay cả khi máy chủ sập nguồn hay mạng ngắt kết nối, quy trình vẫn tự động khôi phục chính xác tại điểm dừng mà không bị mất dữ liệu hay chạy lặp lại các bước đã hoàn tất."
    )
    add_image_block("03_workflow_studio.png", "[Hình 3] Bàn làm việc Visual Workflow Studio - Quy trình cứu khách khiếu nại (WF-CSKH-RECOVERY) với phân luồng rẽ nhánh điều kiện")

    # PAIN POINT 3
    doc.add_heading(level=2).add_run("PAIN POINT 3: Doanh nghiệp e ngại rủi ro khi AI tự ý ra quyết định gây thiệt hại tài chính hoặc sai lệch phát ngôn.")
    p = doc.add_paragraph()
    p.add_run("• Vấn đề thực tế: ").bold = True
    p.add_run("Ảo giác (Hallucination) và thiếu cơ chế rào chắn là rào cản lớn nhất ngăn doanh nghiệp áp dụng AI vào vận hành thực tế. Nếu trao quyền tự động hoàn toàn cho AI, doanh nghiệp có nguy cơ đối mặt với việc AI tự ý cấp phát ngân sách sai, cấp mã giảm giá vượt hạn mức cho phép, hoặc đăng tải các thông điệp truyền thông chưa được kiểm chứng lên mạng xã hội.")

    p = doc.add_paragraph()
    p.add_run("• Giải pháp trong OpenDX CompanyOS:").bold = True
    doc.add_paragraph(
        "1. Mô hình Cổng Phê Duyệt Rủi Ro (Human-in-the-Loop): Hệ thống thiết lập ranh giới an toàn nghiêm ngặt. AI chỉ đóng vai trò phân tích dữ liệu, tổng hợp tình huống và đề xuất phương án tối ưu. Đối với các hành động vượt ngưỡng rủi ro, hệ thống bắt buộc phải dừng lại tại Cổng phê duyệt (Approval Gate) để người có thẩm quyền ký xác nhận.\n"
        "2. Approval Inbox tập trung: Cung cấp giao diện trực quan cho người duyệt kiểm tra đầy đủ ngữ cảnh: Nguyên nhân phát sinh, mức độ thiệt hại của khách hàng, lý do AI đề xuất giải pháp, và nội dung dự thảo phản hồi. Người duyệt có toàn quyền chấp thuận, từ chối hoặc chỉnh sửa phương án trực tiếp.\n"
        "3. Thông báo tự động theo chức vụ qua Email thực tế: Tích hợp giao thức SMTP/Gmail để gửi thông báo tức thời đến đích danh hòm thư của từng vị trí phụ trách (Trưởng phòng Marketing, Trưởng phòng CSKH, Giám đốc Điều hành) kèm đường dẫn xử lý nhanh.\n"
        "4. Nhật ký kiểm toán bất biến (Audit Trail): Mọi thao tác phê duyệt, từ chối, danh tính người xử lý, thời gian thực hiện và lý do đều được ghi nhận vào nhật ký kiểm toán phục vụ hậu kiểm."
    )
    add_image_block("05_approval_inbox.png", "[Hình 4] Hòm thư Phê Duyệt Kiểm Soát Rủi Ro (Approval Inbox) - Thẩm định các đề xuất bồi thường và nội dung chiến dịch")

    # PAIN POINT 4
    doc.add_heading(level=2).add_run("PAIN POINT 4: Dữ liệu vận hành bị phân mảnh và chi phí tiêu thụ tài nguyên AI là một 'hộp đen' khó đo lường.")
    p = doc.add_paragraph()
    p.add_run("• Vấn đề thực tế: ").bold = True
    p.add_run("1. Dữ liệu giữa các khâu bán hàng, tồn kho, thanh toán và hỗ trợ sau bán hàng thường nằm ở các hệ thống riêng lẻ, khiến người quản trị không nắm được bức tranh tổng thể về hoạt động kinh doanh và dịch vụ khách hàng.\n2. Việc sử dụng các mô hình ngôn ngữ lớn (LLM API) thường phát sinh chi phí khó kiểm soát. Doanh nghiệp không có công cụ theo dõi lượng token tiêu thụ theo thời gian thực và không biết được chi phí đó quy đổi ra tiền thực tế là bao nhiêu để cân đối hiệu quả chi phí (ROI).")

    p = doc.add_paragraph()
    p.add_run("• Giải pháp trong OpenDX CompanyOS:").bold = True
    doc.add_paragraph(
        "1. Backend Authoritative Truth (Chân lý dữ liệu từ Backend): Tuân thủ Clean Architecture nghiêm ngặt. Mọi dữ liệu về giá bán, chiết khấu, tồn kho khả dụng và trạng thái thanh toán đều được tính toán và kiểm chứng ở tầng nghiệp vụ lõi (Domain/Application layer). Phía giao diện (Storefront/Console) chỉ hiển thị dữ liệu phản ánh từ backend, loại bỏ hoàn toàn nguy cơ gian lận từ phía client.\n"
        "2. Hồ sơ khách hàng hợp nhất (Customer 360 CRM): Liên kết đồng bộ danh tính khách hàng với toàn bộ lịch sử đơn hàng, phiếu hỗ trợ kỹ thuật, ghi chú chăm sóc và giá trị đóng góp vòng đời (LTV).\n"
        "3. Giám sát chi phí AI minh bạch (AI Cost & Token Observability): Hệ thống tự động ghi nhận số lượng input token, output token, mã định danh mô hình và chi phí vi mô (micros) của từng lần gọi AI vào bảng lưu vết agentic_model_runs. Trên Dashboard điều hành, hệ thống tổng hợp và quy đổi trực tiếp lượng token tiêu thụ ra tiền tệ (VNĐ và USD) theo thời gian thực, giúp người quản trị nắm bắt chính xác mức tiêu hao ngân sách của hệ thống AI."
    )
    add_image_block("01_dashboard.png", "[Hình 5] Bảng Điều Hành Doanh Nghiệp (Commerce Dashboard) - Đo lường tài chính mô phỏng và Chỉ số Token AI quy đổi ra tiền")
    add_image_block("02_company_overview.png", "[Hình 6] Bảng Tổng Quan Vận Hành Công Ty (Company Overview) - Trạng thái hoạt động 4 trụ cột và lộ trình kiến trúc")
    add_image_block("07_customers_crm.png", "[Hình 7] Phân hệ Quản Trị Khách Hàng CRM 360 - Quản lý danh sách khách hàng và lịch sử tương tác")

    # 3. CHẤT LƯỢNG KỸ THUẬT & TÍNH SẴN SÀNG NGUỒN MỞ
    h3 = doc.add_heading(level=1)
    run_h3 = h3.add_run("3. CHẤT LƯỢNG KỸ THUẬT & TÍNH SẴN SÀNG NGUỒN MỞ")
    run_h3.font.name = "Arial"
    run_h3.font.size = Pt(13.5)
    run_h3.font.bold = True
    run_h3.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "1. Kiến trúc phân tầng chuẩn mực (Clean Architecture): Tách biệt hoàn toàn giữa tầng Domain, Application, Infrastructure và Presentation. Logic nghiệp vụ độc lập với framework, cơ sở dữ liệu và giao diện người dùng, giúp hệ thống dễ dàng bảo trì, mở rộng và kiểm thử độc lập.\n"
        "2. Độ tin cậy mã nguồn và tự động hóa kiểm thử: Toàn bộ mã nguồn vượt qua 100% các bài kiểm tra tự động: Lệnh kiểm thử pnpm check chạy qua 23 tệp kiểm thử với 119/119 unit & integration tests đều đạt kết quả Pass. Lệnh pnpm audit:repo đạt chuẩn toàn vẹn cấu trúc mã nguồn mở. Hệ thống kiểm tra kiểu dữ liệu tĩnh TypeScript đạt mức tuyệt đối không có lỗi (0 errors).\n"
        "3. Tuân thủ tiêu chí OLP 2026: Đầy đủ tiêu đề bản quyền SPDX Apache-2.0 trên từng tệp tin nguồn. Triển khai hoàn chỉnh qua Docker Compose, cho phép dựng toàn bộ hệ thống từ mã nguồn gốc mà không phụ thuộc vào các gói cài đặt đóng kín ngoài repo. Bảo đảm nguyên tắc bảo vệ quyền riêng tư dữ liệu (PII-free) trong các tầng báo cáo điều hành."
    )

    # 4. KẾT LUẬN
    h4 = doc.add_heading(level=1)
    run_h4 = h4.add_run("4. KẾT LUẬN")
    run_h4.font.name = "Arial"
    run_h4.font.size = Pt(13.5)
    run_h4.font.bold = True
    run_h4.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "OpenDX CompanyOS giải quyết bài toán cốt lõi của doanh nghiệp khi ứng dụng trí tuệ nhân tạo: Đưa AI vào khuôn khổ vận hành có tổ chức. Bằng việc kết hợp chặt chẽ giữa khả năng lập luận của AI, tính bền bỉ của quy trình thị giác và cơ chế kiểm soát rủi ro của con người, sản phẩm cung cấp một kiến trúc hoàn chỉnh, thực tế và sẵn sàng cho việc mở rộng trong tương lai."
    )

    out_path = os.path.abspath("docs/Thuyet-Minh-SanPham-OpenDX-CompanyOS.docx")
    doc.save(out_path)
    print(f"Document successfully created: {out_path} ({os.path.getsize(out_path) / 1024:.1f} KB)")

if __name__ == "__main__":
    create_document()
